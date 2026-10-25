"""Openpyxl renderers for the legacy payroll workbook downloads."""

from __future__ import annotations

from copy import copy
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.cell import get_column_letter

_CURRENCY_FORMAT = "#,##0.00 [$₽-ru-RU]"
_HEADER_FILL = PatternFill(fill_type="solid", fgColor="FFE2E8F0")


def _workbook() -> Workbook:
    workbook = Workbook()
    workbook.properties.creator = "Consilium"
    return workbook


def _bytes(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _set_controlled_font(sheet: Any) -> None:
    for row in sheet.iter_rows():
        for cell in row:
            font = copy(cell.font)
            font.name = "Arial"
            cell.font = font


def build_payroll_workbook(year: int, month: int, rows: list[dict[str, Any]]) -> bytes:
    workbook = _workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Платёжная ведомость"
    sheet.freeze_panes = "A4"
    sheet.append(["Платёжная ведомость"])
    sheet.append([f"{month:02d}.{year}"])
    headers = [
        "Табельный номер",
        "Сотрудник",
        "Подразделение",
        "Должность",
        "Ставка",
        "План",
        "Начислено",
        "НДФЛ",
        "К выплате",
        "Взносы работодателя",
        "Стоимость работодателя",
        "Отклонение",
    ]
    sheet.append(headers)
    for row in rows:
        sheet.append(
            [
                row["code"],
                row["fullName"],
                row["department"],
                row["position"],
                row["rate"],
                row["plannedGross"],
                row["gross"],
                row["tax"],
                row["payable"],
                row["contributions"],
                row["employerCost"],
                row["variance"],
            ]
        )
    sheet["A1"].font = Font(name="Arial", bold=True, size=16)
    for cell in sheet[3]:
        cell.font = Font(name="Arial", bold=True)
        cell.fill = _HEADER_FILL
    widths = (16, 34, 24, 28, 10, 16, 16, 16, 16, 22, 22, 16)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row_index in range(4, sheet.max_row + 1):
        for column_index in range(6, 13):
            sheet.cell(row_index, column_index).number_format = _CURRENCY_FORMAT
    sheet.auto_filter.ref = "A3:L3"
    _set_controlled_font(sheet)
    return _bytes(workbook)


def build_payslip_workbook(year: int, month: int, row: dict[str, Any]) -> bytes:
    workbook = _workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Расчетный лист"
    sheet.append(["РАСЧЕТНЫЙ ЛИСТ"])
    sheet.append([f"За период: {month:02d}.{year}"])
    sheet.append([])
    sheet.append(["Табельный номер:", row["code"]])
    sheet.append(["Сотрудник:", row["fullName"]])
    sheet.append(["Подразделение:", row["department"]])
    sheet.append(["Должность:", row["position"]])
    sheet.append(["Ставка:", row["rate"]])
    sheet.append([])
    sheet.append(["Вид начисления/удержания", "План", "Факт"])
    sheet.append(["Оклад", row["plannedGross"], row["actualBase"]])
    sheet.append(["Надбавки", 0, row["allowance"]])
    sheet.append(["Премии", 0, row["bonus"] + row["oneTime"]])
    sheet.append(["Удержания (НДФЛ)", 0, -row["tax"]])
    sheet.append(["Прочие удержания", 0, -row["deduction"]])
    sheet.append([])
    sheet.append(["Всего начислено:", "", row["gross"]])
    sheet.append(["НДФЛ удержан:", "", -row["tax"]])
    sheet.append(["К выплате на руки:", "", row["payable"]])
    sheet.append(["Взносы ПФР/ФОМС/ФСС:", "", row["contributions"]])
    sheet.append([])
    sheet["A1"].font = Font(name="Arial", bold=True, size=14)
    for cell in sheet[10]:
        cell.font = Font(name="Arial", bold=True)
        cell.fill = _HEADER_FILL
    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 15
    sheet.column_dimensions["C"].width = 15
    for row_index in range(11, 21):
        sheet.cell(row_index, 2).number_format = _CURRENCY_FORMAT
        sheet.cell(row_index, 3).number_format = _CURRENCY_FORMAT
    _set_controlled_font(sheet)
    return _bytes(workbook)


class OpenpyxlPayrollWorkbook:
    """Openpyxl adapter for the application payroll workbook port."""

    def build_payroll(self, year: int, month: int, rows: list[dict[str, Any]]) -> bytes:
        return build_payroll_workbook(year, month, rows)

    def build_payslip(self, year: int, month: int, row: dict[str, Any]) -> bytes:
        return build_payslip_workbook(year, month, row)
