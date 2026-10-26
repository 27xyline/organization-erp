"""Синхронные средства формирования экспорта; вызывающий код запускает их \
в рабочем потоке после UoW базы данных."""

from __future__ import annotations

import asyncio
import re
from io import BytesIO
from pathlib import Path
from threading import Lock
from typing import Any

import qrcode  # type: ignore[import-untyped]
import reportlab
from docx import Document as WordDocument
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.cell import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

_FONT_REGISTRATION_LOCK = Lock()


def safe_filename(value: str, extension: str) -> str:
    stem = re.sub(r"[^\w.-]+", "_", value, flags=re.UNICODE).strip("._") or "report"
    return f"{stem[:100]}.{extension.lstrip('.')}"


def render_xlsx(title: str, columns: list[str], rows: list[list[Any]]) -> bytes:
    workbook = Workbook()
    workbook.properties.creator = "Consilium"
    sheet = workbook.active
    assert sheet is not None
    sheet.title = title[:31] or "Отчёт"
    sheet.append(columns)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(name="Arial", bold=True, color="FF0F172A")
        cell.fill = PatternFill(fill_type="solid", fgColor="FFE2E8F0")
    for column in range(1, sheet.max_column + 1):
        width = max(
            len(str(sheet.cell(row, column).value or "")) for row in range(1, sheet.max_row + 1)
        )
        sheet.column_dimensions[get_column_letter(column)].width = min(40, max(14, width + 2))
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def render_dashboard_xlsx(data: dict[str, Any]) -> bytes:
    """Создаёт книгу панели мониторинга из уже авторизованного отделённого DTO."""
    workbook = Workbook()
    workbook.properties.creator = "Consilium"
    overview = workbook.active
    assert overview is not None
    overview.title = "Сводка"
    overview.append(["Показатель", "Значение"])
    overview.append(["Период", f"{data['query']['dateFrom']} — {data['query']['dateTo']}"])
    labels = (
        ("employees", "Сотрудников", "headcount"),
        ("employees", "Активных сотрудников", "active"),
        ("employees", "Занято ставок", "occupiedRate"),
        ("employees", "Штатных ставок", "staffRate"),
        ("finance", "Плановый ФОТ", "plannedFot"),
        ("finance", "Фактический ФОТ", "actualFot"),
        ("employees", "Занятость штата, %", "occupancy"),
        ("projects", "Активных проектов", "count"),
        ("projects", "Плановый бюджет", "plannedBudget"),
        ("projects", "Фактический бюджет", "actualBudget"),
        ("projects", "Средний прогресс, %", "averageProgress"),
        ("assets", "Стоимость имущества", "assetValue"),
        ("tasks", "Просроченных задач", "overdueTasks"),
    )
    for scope, title, key in labels:
        if scope == "tasks" or data["visibility"].get(scope, False):
            overview.append([title, data["summary"].get(key, 0)])
    _style_worksheet(overview)
    workforce_headers = [
        "Подразделение",
        "Сотрудников",
        "Активных",
        "Занято ставок",
        "Штатных ставок",
        "Занятость, %",
    ]
    if data["visibility"].get("finance"):
        workforce_headers.extend(["Плановый ФОТ", "Фактический ФОТ"])
    workforce_rows = []
    if data["visibility"].get("employees"):
        for row in data["analytics"]["workforceByDepartment"]:
            item = [
                row["department"],
                row["headcount"],
                row["active"],
                row["occupiedRate"],
                row["staffRate"],
                row["occupancy"],
            ]
            if data["visibility"].get("finance"):
                item.extend([row["plannedFot"], row["actualFot"]])
            workforce_rows.append(item)
    _append_sheet(
        workbook,
        "Сотрудники и ФОТ" if data["visibility"].get("finance") else "Сотрудники",
        workforce_headers,
        workforce_rows,
    )
    _append_sheet(
        workbook,
        "Проекты",
        [
            "Код",
            "Наименование",
            "Окончание",
            "Прогресс, %",
            "Задач",
            "Выполнено",
            "Плановый бюджет",
            "Фактический бюджет",
            "Отклонение",
        ],
        [
            [
                r["code"],
                r["name"],
                r["endDate"],
                r["progress"],
                r["tasksTotal"],
                r["tasksCompleted"],
                r["plannedBudget"],
                r["actualBudget"],
                r["actualBudget"] - r["plannedBudget"],
            ]
            for r in data["activeProjects"]
        ]
        if data["visibility"].get("projects")
        else [],
    )
    _append_sheet(
        workbook,
        "Имущество",
        ["Статус", "Количество", "Стоимость"],
        [[r["status"], r["count"], r["value"]] for r in data["analytics"]["assetsByStatus"]]
        if data["visibility"].get("assets")
        else [],
    )
    _append_sheet(
        workbook,
        "Отпуска",
        ["Сотрудник", "Подразделение", "Тип", "Начало", "Окончание"],
        [
            [r["employee"], r["department"], r["type"], r["startDate"], r["endDate"]]
            for r in data["analytics"]["vacations"]
        ],
    )
    _append_sheet(
        workbook,
        "Контроль сроков",
        ["Категория", "Объект", "Контекст", "Срок / статус"],
        [
            *[
                [
                    "Просроченная задача",
                    r["name"],
                    f"{r['project']['code']} · {r['responsible']}",
                    f"{r['overdueDays']} дн.",
                ]
                for r in data["overdueTasks"]
            ],
            *[
                ["Трудовой договор", r["employee"], r["department"], r["contractEndDate"]]
                for r in data["contracts"]
            ],
            *[
                ["Имущество", r["name"], r["inventoryNumber"], r["status"]]
                for r in data["assetAttention"]
            ],
        ],
    )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


class DashboardExportRenderer:
    """Асинхронный адаптер интерфейса; вся работа с документами, нагружающая CPU, \
    выполняется вне цикла событий."""

    async def render(self, output_format: str, data: dict[str, Any]) -> bytes:
        if output_format == "xlsx":
            return await asyncio.to_thread(render_dashboard_xlsx, data)
        columns, rows = _dashboard_report_rows(data)
        renderer = render_docx if output_format == "docx" else render_pdf
        return await asyncio.to_thread(renderer, "Отчёт по dashboard", columns, rows)


def _dashboard_report_rows(data: dict[str, Any]) -> tuple[list[str], list[list[Any]]]:
    """Возвращает подробные таблицы в заданной области в плоском формате для DOCX и PDF."""
    rows: list[list[Any]] = []

    def add_section(title: str, headers: list[str], items: list[list[Any]]) -> None:
        rows.append([title, "", "", ""])
        for item in items:
            rows.extend(
                [title, str(item[0]), header, value]
                for header, value in zip(headers[1:], item[1:], strict=True)
            )

    summary = data["summary"]
    rows.append(["Сводка", "", "", ""])
    rows.append(
        ["Сводка", "Период", "", f"{data['query']['dateFrom']} — {data['query']['dateTo']}"]
    )
    for scope, label, key in (
        ("employees", "Сотрудников", "headcount"),
        ("employees", "Активных сотрудников", "active"),
        ("employees", "Занято ставок", "occupiedRate"),
        ("employees", "Штатных ставок", "staffRate"),
        ("finance", "Плановый ФОТ", "plannedFot"),
        ("finance", "Фактический ФОТ", "actualFot"),
        ("projects", "Активных проектов", "count"),
        ("projects", "Плановый бюджет", "plannedBudget"),
        ("projects", "Фактический бюджет", "actualBudget"),
        ("assets", "Стоимость имущества", "assetValue"),
    ):
        if data["visibility"].get(scope, False):
            rows.append(["Сводка", "", label, summary.get(key, 0)])

    if data["visibility"].get("employees"):
        workforce_headers = [
            "Подразделение",
            "Сотрудников",
            "Активных",
            "Занято ставок",
            "Штатных ставок",
            "Занятость, %",
        ]
        if data["visibility"].get("finance"):
            workforce_headers.extend(["Плановый ФОТ", "Фактический ФОТ"])
        add_section(
            "Сотрудники и ФОТ" if data["visibility"].get("finance") else "Сотрудники",
            workforce_headers,
            [
                [
                    r["department"],
                    r["headcount"],
                    r["active"],
                    r["occupiedRate"],
                    r["staffRate"],
                    r["occupancy"],
                ]
                + ([r["plannedFot"], r["actualFot"]] if data["visibility"].get("finance") else [])
                for r in data["analytics"]["workforceByDepartment"]
            ],
        )
    if data["visibility"].get("projects"):
        add_section(
            "Проекты",
            [
                "Код",
                "Наименование",
                "Окончание",
                "Прогресс, %",
                "Задач",
                "Выполнено",
                "Плановый бюджет",
                "Фактический бюджет",
                "Отклонение",
            ],
            [
                [
                    r["code"],
                    r["name"],
                    r["endDate"],
                    r["progress"],
                    r["tasksTotal"],
                    r["tasksCompleted"],
                    r["plannedBudget"],
                    r["actualBudget"],
                    r["actualBudget"] - r["plannedBudget"],
                ]
                for r in data["activeProjects"]
            ],
        )
    if data["visibility"].get("assets"):
        add_section(
            "Имущество",
            ["Статус", "Количество", "Стоимость"],
            [[r["status"], r["count"], r["value"]] for r in data["analytics"]["assetsByStatus"]],
        )
    add_section(
        "Отпуска",
        ["Сотрудник", "Подразделение", "Тип", "Начало", "Окончание"],
        [
            [r["employee"], r["department"], r["type"], r["startDate"], r["endDate"]]
            for r in data["analytics"]["vacations"]
        ],
    )
    add_section(
        "Контроль сроков",
        ["Категория", "Объект", "Контекст", "Срок / статус"],
        [
            *[
                [
                    "Просроченная задача",
                    r["name"],
                    f"{r['project']['code']} · {r['responsible']}",
                    f"{r['overdueDays']} дн.",
                ]
                for r in data["overdueTasks"]
            ],
            *[
                ["Трудовой договор", r["employee"], r["department"], r["contractEndDate"]]
                for r in data["contracts"]
            ],
            *[
                ["Имущество", r["name"], r["inventoryNumber"], r["status"]]
                for r in data["assetAttention"]
            ],
        ],
    )
    return ["Раздел", "Объект", "Показатель", "Значение"], rows


def _style_worksheet(sheet: Any) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(name="Arial", bold=True, color="FF0F172A")
        cell.fill = PatternFill(fill_type="solid", fgColor="FFE2E8F0")
    for column in range(1, sheet.max_column + 1):
        width = max(
            len(str(sheet.cell(row, column).value or "")) for row in range(1, sheet.max_row + 1)
        )
        sheet.column_dimensions[get_column_letter(column)].width = min(40, max(14, width + 2))


def _append_sheet(
    workbook: Workbook, title: str, headers: list[str], rows: list[list[Any]]
) -> None:
    sheet = workbook.create_sheet(title=title)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    _style_worksheet(sheet)


def render_docx(title: str, columns: list[str], rows: list[list[Any]]) -> bytes:
    document = WordDocument()
    document.add_heading(title, level=1)
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, columns, strict=True):
        cell.text = str(value)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            cell.text = "" if value is None else str(value)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _cyrillic_font(text: str) -> Path:
    assert reportlab.__file__ is not None
    candidates = (
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
        Path(reportlab.__file__).parent / "fonts" / "Vera.ttf",
    )
    for candidate in candidates:
        if not candidate.is_file():
            continue
        font = TTFont("ExportFontProbe", str(candidate))
        if all(ord(character) in font.face.charToGlyph for character in text):
            return candidate
    raise RuntimeError("No installed TrueType font supports every character in this PDF")


def render_pdf(title: str, columns: list[str], rows: list[list[Any]]) -> bytes:
    font_text = title + "".join(str(value) for value in columns)
    font_text += "".join(str(value) for row in rows for value in row if value is not None)
    font_path = _cyrillic_font(font_text)
    with _FONT_REGISTRATION_LOCK:
        if "ConsiliumExport" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("ConsiliumExport", str(font_path)))
    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=landscape(letter), title=title)
    data = [columns] + [["" if value is None else str(value) for value in row] for row in rows]
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), "ConsiliumExport"),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0")),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#94a3b8")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("ExportTitle", parent=styles["Title"], fontName="ConsiliumExport")
    document.build([Paragraph(title, title_style), table])
    return output.getvalue()


def render_qr(value: str) -> bytes:
    code = qrcode.QRCode(error_correction=qrcode.constants.ERROR_CORRECT_Q, border=2)
    code.add_data(value)
    code.make(fit=True)
    output = BytesIO()
    code.make_image().save(output, format="PNG")
    return output.getvalue()
