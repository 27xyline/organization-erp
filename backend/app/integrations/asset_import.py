"""Адаптеры Openpyxl для старых книг импорта, шаблонов и экспорта основных средств."""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from app.domain.assets.importing import (
    ASSET_IMPORT_COLUMNS,
    MAX_ASSET_IMPORT_ROWS,
    AssetWorkbookFormatError,
    ParsedAssetImportRow,
)


def _append_values(sheet: Any, values: Sequence[object]) -> None:
    sheet.append(list(values))
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str) and cell.value.startswith("="):
            cell.data_type = "s"


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value).lower()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _cell_value(value: object) -> object:
    if value is None:
        return ""
    if isinstance(value, (datetime, int, float, str, bool)):
        return value
    return str(value)


def parse_asset_import_workbook(payload: bytes) -> tuple[ParsedAssetImportRow, ...]:
    """Читает ячейки книги и сопоставляет известные заголовки; \
    бизнес-проверка остаётся в сервисах."""
    try:
        workbook = load_workbook(BytesIO(payload), data_only=True, read_only=False)
    except Exception as error:
        raise AssetWorkbookFormatError(
            "Не удалось прочитать файл. Загрузите корректный шаблон .xlsx"
        ) from error

    sheet = workbook["Имущество"] if "Имущество" in workbook.sheetnames else None
    if (
        sheet is None
        or sheet.max_row < 1
        or (sheet.max_row == 1 and sheet.max_column == 1 and sheet.cell(1, 1).value is None)
    ):
        raise AssetWorkbookFormatError("В книге нет листа «Имущество»")
    if sheet.max_row > MAX_ASSET_IMPORT_ROWS + 1:
        raise AssetWorkbookFormatError(f"В файле больше {MAX_ASSET_IMPORT_ROWS} строк с данными")

    known_by_label = {column.label: column for column in ASSET_IMPORT_COLUMNS}
    seen_labels: set[str] = set()
    column_by_index: dict[int, str] = {}
    unknown_labels: list[str] = []
    for index, cell in enumerate(sheet[1], start=1):
        label = _cell_text(_cell_value(cell.value))
        if not label:
            continue
        if label in seen_labels:
            raise AssetWorkbookFormatError(f"Колонка «{label}» указана дважды")
        seen_labels.add(label)
        column = known_by_label.get(label)
        if column is None:
            unknown_labels.append(label)
        else:
            column_by_index[index] = column.key

    if unknown_labels:
        raise AssetWorkbookFormatError(f"Неизвестные колонки: {', '.join(unknown_labels)}")
    missing = [
        column.label
        for column in ASSET_IMPORT_COLUMNS
        if column.required and column.key not in column_by_index.values()
    ]
    if missing:
        raise AssetWorkbookFormatError(f"Не найдены обязательные колонки: {', '.join(missing)}")

    rows: list[ParsedAssetImportRow] = []
    for row_number in range(2, sheet.max_row + 1):
        if not any(
            _cell_text(_cell_value(sheet.cell(row=row_number, column=index).value))
            for index in range(1, len(ASSET_IMPORT_COLUMNS) + 1)
        ):
            continue
        values = {
            key: _cell_value(sheet.cell(row=row_number, column=index).value)
            for index, key in column_by_index.items()
        }
        rows.append(ParsedAssetImportRow(row_number, values))

    if not rows:
        raise AssetWorkbookFormatError("В файле нет строк с имуществом")
    return tuple(rows)


def build_asset_import_template(
    mols: Sequence[Mapping[str, Any]], groups: Sequence[Mapping[str, Any]]
) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])
    workbook.properties.creator = "Consilium"
    workbook.properties.created = datetime.now()

    guide = workbook.create_sheet("Как заполнить")
    guide.column_dimensions["A"].width = 100
    guide.append(["Импорт имущества"])
    guide.append(["Заполните лист «Имущество» и загрузите этот файл в разделе имущества."])
    guide.append(
        [f"Допустимо до {MAX_ASSET_IMPORT_ROWS} строк за один импорт. Пустые строки пропускаются."]
    )
    guide.append(
        [
            "Все строки должны пройти проверку до записи. Если есть ошибка или дубликат, "
            "имущество не будет создано."
        ]
    )
    guide.append(["Коды МОЛ и групп выбирайте из списков или копируйте с листа «Справочники»."])
    guide.append(
        [
            "Цена и количество должны быть положительными и содержать не более двух знаков "
            "после запятой."
        ]
    )
    guide.append(["Дату можно указать как дату Excel или в формате ДД.ММ.ГГГГ."])
    guide.append(["Ранее учтённое имущество укажите как «Да» или «Нет». Форма учета: 145 или 367."])
    guide.append(
        [
            "После загрузки проверьте предварительный просмотр: он показывает ошибки и "
            "повторяющиеся номера."
        ]
    )
    guide["A1"].font = Font(bold=True, size=16, color="FF12304A")
    for row in guide.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="center")
        guide.row_dimensions[row[0].row].height = 30

    sheet = workbook.create_sheet("Имущество")
    for index, column in enumerate(ASSET_IMPORT_COLUMNS, start=1):
        sheet.cell(row=1, column=index, value=column.label)
        sheet.column_dimensions[get_column_letter(index)].width = max(
            16, min(34, len(column.label) + 4)
        )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:O{MAX_ASSET_IMPORT_ROWS + 1}"
    sheet.row_dimensions[1].height = 32
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="FF12304A")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row_number in range(2, MAX_ASSET_IMPORT_ROWS + 2):
        sheet.cell(row=row_number, column=3).number_format = "#,##0.00"
        sheet.cell(row=row_number, column=5).number_format = "#,##0.00"
        sheet.cell(row=row_number, column=8).number_format = "dd.mm.yyyy"

    references = workbook.create_sheet("Справочники")
    for index, (header, width) in enumerate(
        (
            ("Код МОЛ", 18),
            ("Ответственный", 36),
            ("Подразделение", 34),
            ("Код группы", 18),
            ("Группа имущества", 34),
        ),
        start=1,
    ):
        references.cell(row=1, column=index, value=header)
        references.column_dimensions[
            references.cell(row=1, column=index).column_letter
        ].width = width
        cell = references.cell(row=1, column=index)
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="FF12304A")
    for index in range(max(len(mols), len(groups))):
        mol = mols[index] if index < len(mols) else {}
        group = groups[index] if index < len(groups) else {}
        _append_values(
            references,
            [
                mol.get("code") or "",
                mol.get("fullName") or "",
                mol.get("department") or "",
                group.get("code") or "",
                group.get("name") or "",
            ],
        )

    if mols:
        workbook.defined_names.add(
            DefinedName("MolCodes", attr_text=f"'Справочники'!$A$2:$A${len(mols) + 1}")
        )
        mol_validation = DataValidation(type="list", allow_blank=False, formula1="=MolCodes")
        sheet.add_data_validation(mol_validation)
        mol_validation.add(f"F2:F{MAX_ASSET_IMPORT_ROWS + 1}")
    if groups:
        workbook.defined_names.add(
            DefinedName("AssetGroupCodes", attr_text=f"'Справочники'!$D$2:$D${len(groups) + 1}")
        )
        group_validation = DataValidation(
            type="list", allow_blank=False, formula1="=AssetGroupCodes"
        )
        sheet.add_data_validation(group_validation)
        group_validation.add(f"G2:G{MAX_ASSET_IMPORT_ROWS + 1}")
    existing_validation = DataValidation(type="list", allow_blank=True, formula1='"Да,Нет"')
    accounting_validation = DataValidation(type="list", allow_blank=True, formula1='"145,367"')
    sheet.add_data_validation(existing_validation)
    sheet.add_data_validation(accounting_validation)
    existing_validation.add(f"M2:M{MAX_ASSET_IMPORT_ROWS + 1}")
    accounting_validation.add(f"N2:N{MAX_ASSET_IMPORT_ROWS + 1}")

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_export_workbook(headers: Sequence[str], rows: Sequence[Mapping[str, object]]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])
    sheet = workbook.create_sheet("Data")
    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    for index, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(14, min(40, len(header) + 4))
    for record in rows:
        _append_values(sheet, [record.get(header, "") for header in headers])
    if headers:
        sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


class OpenpyxlAssetWorkbook:
    """Адаптер Openpyxl, реализующий доменный интерфейс для книг."""

    def parse_workbook(self, payload: bytes) -> tuple[ParsedAssetImportRow, ...]:
        return parse_asset_import_workbook(payload)

    def build_import_template(
        self,
        mols: Sequence[Mapping[str, object]],
        groups: Sequence[Mapping[str, object]],
    ) -> bytes:
        return build_asset_import_template(mols, groups)

    def build_export(
        self,
        headers: Sequence[str],
        rows: Sequence[Mapping[str, object]],
    ) -> bytes:
        return build_export_workbook(headers, rows)
