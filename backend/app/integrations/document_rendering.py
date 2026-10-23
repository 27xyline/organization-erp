"""Existing order/acceptance templates, rendered outside the event loop."""

import asyncio
from io import BytesIO
from pathlib import Path
from threading import Lock
from typing import Literal
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from app.domain.documents.records import RenderTemplate

FONT_LOCK = Lock()
MONTHS = (
    "января",
    "февраля",
    "марта",
    "апреля",
    "мая",
    "июня",
    "июля",
    "августа",
    "сентября",
    "октября",
    "ноября",
    "декабря",
)
FONT_PATHS = (
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
)


class DocumentRenderer:
    async def render(self, command: RenderTemplate) -> bytes:
        return await asyncio.to_thread(self._render, command)

    def _render(self, command: RenderTemplate) -> bytes:
        date = command.date
        stamp = f"{date.day:02d} {MONTHS[date.month - 1]} {date.year} г."
        order = command.template == "PERSONNEL_ORDER"
        heading = "ПРИКАЗ" if order else "АКТ"
        signatures = (
            ("Руководитель: ____________________", "С приказом ознакомлен(а): ____________________")
            if order
            else ("Передал: ____________________", "Принял: ____________________")
        )
        if command.format == "pdf":
            with FONT_LOCK:
                if "ERPRussian" not in pdfmetrics.getRegisteredFontNames():
                    font = next((path for path in FONT_PATHS if Path(path).is_file()), None)
                    if font is None:
                        raise RuntimeError(
                            "A Cyrillic TrueType font is required for document PDF rendering"
                        )
                    pdfmetrics.registerFont(TTFont("ERPRussian", font))
            output = BytesIO()
            pdf = SimpleDocTemplate(
                output, leftMargin=50, rightMargin=50, topMargin=50, bottomMargin=50
            )

            def paragraph(
                text: str, size: int = 11, alignment: Literal[0, 1, 2, 4] = 0
            ) -> Paragraph:
                return Paragraph(
                    escape(text).replace("\n", "<br/>"),
                    ParagraphStyle(
                        "Document",
                        fontName="ERPRussian",
                        fontSize=size,
                        leading=size + 4,
                        alignment=alignment,
                    ),
                )

            story = [
                paragraph(heading, 20, TA_CENTER),
                Spacer(1, 6),
                paragraph(f"№ {command.number} от {stamp}", 12, TA_CENTER),
                Spacer(1, 22),
                paragraph(command.subject.upper(), 13),
                Spacer(1, 14),
                paragraph(command.details, 11, TA_JUSTIFY),
                Spacer(1, 22),
            ]
            if command.basis:
                story.extend([paragraph("Основание: " + command.basis, 10), Spacer(1, 22)])
            story.extend(
                [Spacer(1, 30), paragraph(signatures[0]), Spacer(1, 14), paragraph(signatures[1])]
            )
            pdf.build(story)
            return output.getvalue()
        doc = Document()
        doc.styles["Normal"].font.name = "Arial"
        doc.styles["Normal"].font.size = Pt(12)

        def line(
            text: str,
            size: int = 12,
            *,
            bold: bool = False,
            center: bool = False,
            justified: bool = False,
            italic: bool = False,
        ) -> None:
            paragraph = doc.add_paragraph()
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if center
                else WD_ALIGN_PARAGRAPH.JUSTIFY
                if justified
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            run = paragraph.add_run(text)
            run.font.size = Pt(size)
            run.bold, run.italic = bold, italic

        line(heading, 16, bold=True, center=True)
        line(f"№ {command.number} от {stamp}", center=True)
        line("")
        line(command.subject.upper(), 13, bold=True)
        line("")
        line(command.details, justified=True)
        line("")
        line("Основание: " + command.basis if command.basis else "", 11, italic=True)
        line("")
        line("")
        line(signatures[0])
        line("")
        line(signatures[1])
        output = BytesIO()
        doc.save(output)
        return output.getvalue()
