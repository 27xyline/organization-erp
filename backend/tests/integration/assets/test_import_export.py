"""Binary workbook contracts and atomic import against the real database."""

from dataclasses import replace
from io import BytesIO

import httpx
import pytest
from openpyxl import Workbook, load_workbook
from opentelemetry import trace
from sqlalchemy import func, select

from app.api.dependencies import get_authenticated_context
from app.domain.assets.importing import ASSET_IMPORT_COLUMNS
from app.integrations.asset_import import OpenpyxlAssetWorkbook
from app.main import create_app
from app.models.assets import Asset
from app.models.users import AuditLog
from app.services.asset_import import AssetImportService


def workbook(prefix, mol_code, group_code, *, duplicate=False):
    book = Workbook()
    sheet = book.active
    sheet.title = "Имущество"
    sheet.append(
        [
            column[1] if isinstance(column, tuple) else column.label
            for column in ASSET_IMPORT_COLUMNS
        ]
    )
    for index in range(2):
        sheet.append(
            [
                prefix + ("-0" if duplicate else "-" + str(index)),
                "Импорт",
                "12,34",
                "шт",
                "0,20",
                mol_code,
                group_code,
                "16.10.2026",
                "Акт",
                "№7",
            ]
        )
    output = BytesIO()
    book.save(output)
    return output.getvalue()


@pytest.mark.anyio
async def test_import_preview_commit_template_and_scoped_binary_export(asset_env):
    prefix, context, factory, service, command, sessions = asset_env
    app = create_app()
    app.state.tracer = trace.get_tracer(__name__)
    app.state.unit_of_work_factory = factory

    async def auth():
        return context

    app.dependency_overrides[get_authenticated_context] = auth
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        template = await client.get("/api/assets/import/template")
        assert template.status_code == 200, template.text
        assert (
            template.headers["content-type"]
            == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert (
            template.headers["content-disposition"]
            == 'attachment; filename="asset-import-template.xlsx"'
        )
        assert template.headers["cache-control"] == "no-store"
        book = load_workbook(BytesIO(template.content))
        assert book.sheetnames == ["Как заполнить", "Имущество", "Справочники"]
        preview = await client.post(
            "/api/assets/import/preview",
            files={
                "file": (
                    "assets.xlsx",
                    workbook(prefix, prefix, prefix),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert preview.status_code == 200, preview.text
        rows = preview.json()["data"]
        assert rows["validCount"] == 2 and rows["errorCount"] == 0, rows
        inputs = [row["input"] for row in rows["rows"]]
        commit = await client.post("/api/assets/import/commit", json={"assets": inputs})
        assert commit.status_code == 201, commit.text
        assert commit.json()["data"] == {
            "importedCount": 2,
            "inventoryNumbers": [prefix + "-0", prefix + "-1"],
        }
        duplicate = await client.post("/api/assets/import/commit", json={"assets": inputs})
        assert duplicate.status_code == 409
        assert duplicate.json()["error"]["code"] == "ASSET_IMPORT_REVIEW_REQUIRED"
        for kind in ["assets", "operations"]:
            exported = await client.get("/api/export?type=" + kind)
            assert exported.status_code == 200, exported.text
            assert (
                exported.headers["content-type"]
                == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            assert f"{kind}_export_" in exported.headers["content-disposition"]
            book = load_workbook(BytesIO(exported.content), data_only=True)
            assert book.sheetnames == ["Data"]
            data = list(book["Data"].values)
            assert any(prefix + "-0" in row for row in data)
        assert (await client.get("/api/export?type=bad")).status_code == 422
        assert (await client.post("/api/assets/import/preview")).json()["error"][
            "code"
        ] == "ASSET_IMPORT_FILE_REQUIRED"
        assert (
            await client.post("/api/assets/import/preview", files={"file": ("bad.csv", b"bad")})
        ).json()["error"]["code"] == "ASSET_IMPORT_FORMAT_UNSUPPORTED"
        assert (
            await client.post("/api/assets/import/preview", files={"file": ("bad.xlsx", b"bad")})
        ).json()["error"]["code"] == "ASSET_IMPORT_INVALID_FILE"
        duplicate_preview = await client.post(
            "/api/assets/import/preview",
            files={"file": ("assets.xlsx", workbook(prefix, prefix, prefix, duplicate=True))},
        )
        assert duplicate_preview.json()["data"]["errorCount"] == 2


@pytest.mark.anyio
async def test_import_second_audit_failure_rolls_back_whole_batch(asset_env):
    prefix, context, factory, service, command, sessions = asset_env

    class AuditFailure:
        def __init__(self, original):
            self.original = original
            self.count = 0

        async def append(self, record):
            self.count += 1
            if self.count == 2:
                raise RuntimeError("injected second audit failure")
            await self.original.append(record)

    def failing_factory():
        uow = factory()
        uow.audit = AuditFailure(uow.audit)
        return uow

    with pytest.raises(RuntimeError, match="injected second audit failure"):
        await AssetImportService(failing_factory, OpenpyxlAssetWorkbook()).commit(
            (
                replace(command, inventoryNumber=prefix + "-0"),
                replace(command, inventoryNumber=prefix + "-1"),
            ),
            context,
        )
    async with sessions() as session:
        assert (
            await session.scalar(
                select(func.count())
                .select_from(Asset)
                .where(Asset.inventoryNumber.like(prefix + "%"))
            )
            == 0
        )
        assert (
            await session.scalar(
                select(func.count()).select_from(AuditLog).where(AuditLog.userId == prefix)
            )
            == 0
        )


@pytest.mark.anyio
async def test_binary_and_inventory_reads_do_not_expose_foreign_department(asset_env):
    from app.domain.auth.rules import AccessContext, ScopeGrant
    from app.services.dto.assets import InventoryCommand
    from app.services.inventory import InventoryService

    prefix, context, factory, service, command, sessions = asset_env
    own = await service.create(command, context)
    foreign = await service.create(
        replace(command, inventoryNumber=prefix + "-foreign", molId=prefix + "-other"), context
    )
    foreign_sheet = await InventoryService(factory).create_inventory(
        InventoryCommand("Чужая ведомость", prefix + "-other"), context
    )
    scoped = replace(
        context,
        actor=replace(
            context.actor,
            access=AccessContext(
                prefix,
                None,
                None,
                (),
                (ScopeGrant("scope", "ASSET_CUSTODIAN", "ASSIGNED", "NONE", (prefix,)),),
            ),
        ),
    )
    app = create_app()
    app.state.tracer = trace.get_tracer(__name__)
    app.state.unit_of_work_factory = factory

    async def auth():
        return scoped

    app.dependency_overrides[get_authenticated_context] = auth
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        assert (await client.get("/api/assets/" + foreign.id)).status_code == 403
        assert (await client.get("/api/assets/" + foreign.id + "/qr")).status_code == 404
        assert (await client.get("/api/assets/" + own.id + "/qr")).status_code == 200
        listing = await client.get("/api/assets?search=" + prefix)
        assert [row["id"] for row in listing.json()["data"]] == [own.id]
        for kind in ("assets", "operations"):
            response = await client.get("/api/export?type=" + kind)
            book = load_workbook(BytesIO(response.content), data_only=True)
            all_values = list(book["Data"].values)
            assert any(prefix in row for row in all_values)
            assert not any(prefix + "-foreign" in row for row in all_values)
        assert (
            await client.get("/api/assets/inventories/" + foreign_sheet["id"])
        ).status_code == 404
        assert (
            await client.get("/api/assets/inventories/" + foreign_sheet["id"] + "/act")
        ).status_code == 404
        assert (
            await client.get("/api/assets/inventories/resolve?number=" + prefix + "-foreign")
        ).json()["data"] == []
        assert (
            await client.post(
                "/api/assets/inventories/" + foreign_sheet["id"] + "/scan",
                json={"inventoryNumber": prefix + "-foreign", "foundQuantity": 0},
            )
        ).status_code == 404
        assert (
            await client.post("/api/assets/inventories/" + foreign_sheet["id"] + "/complete")
        ).status_code == 404
        options = (await client.get("/api/assets/inventories/options")).json()["data"]
        assert all(row["departmentId"] == prefix for row in options)
        template = await client.get("/api/assets/import/template")
        book = load_workbook(BytesIO(template.content), data_only=True)
        assert not any(prefix + "-other" in row for row in book["Справочники"].values)
