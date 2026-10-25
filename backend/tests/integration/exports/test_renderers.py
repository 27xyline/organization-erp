from io import BytesIO
from typing import Any

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from app.integrations.exports import (
    _dashboard_report_rows,
    render_dashboard_xlsx,
    render_docx,
    render_pdf,
    render_qr,
    render_xlsx,
    safe_filename,
)


def test_renderers_keep_russian_labels_and_numeric_values_readable() -> None:
    columns = ["Сотрудник", "Начислено"]
    rows = [["Иван Иванов", 1250.5]]

    xlsx = load_workbook(BytesIO(render_xlsx("Отчёт", columns, rows)), data_only=True)
    assert xlsx.active is not None
    assert list(xlsx.active.values) == [("Сотрудник", "Начислено"), ("Иван Иванов", 1250.5)]

    docx = Document(BytesIO(render_docx("Отчёт", columns, rows)))
    assert docx.tables[0].cell(1, 0).text == "Иван Иванов"
    assert docx.tables[0].cell(1, 1).text == "1250.5"

    pdf = render_pdf("Отчёт", columns, rows)
    assert pdf.startswith(b"%PDF-")
    assert b"/FontFile2" in pdf  # встроенный шрифт TrueType обеспечивает отображение кириллицы

    qr = render_qr("https://erp.example/документы")
    assert qr.startswith(b"\x89PNG\r\n\x1a\n")


def test_filename_is_safe_and_keeps_russian_letters() -> None:
    assert safe_filename("Отчёт: октябрь / 2026", "xlsx") == "Отчёт_октябрь_2026.xlsx"


def test_docx_and_pdf_exports_include_scoped_detail_sections() -> None:
    data = {
        "query": {"dateFrom": "2026-01-01", "dateTo": "2026-12-31"},
        "visibility": {"employees": True, "finance": True, "projects": True, "assets": True},
        "summary": {
            "headcount": 3,
            "active": 2,
            "occupiedRate": 2.5,
            "staffRate": 3,
            "plannedFot": 1000,
            "actualFot": 900,
            "count": 1,
            "plannedBudget": 5000,
            "actualBudget": 4500,
            "assetValue": 300,
        },
        "analytics": {
            "workforceByDepartment": [
                {
                    "department": "Отдел разработки",
                    "headcount": 3,
                    "active": 2,
                    "occupiedRate": 2.5,
                    "staffRate": 3,
                    "occupancy": 83.3,
                    "plannedFot": 1000,
                    "actualFot": 900,
                }
            ],
            "assetsByStatus": [{"status": "ACTIVE", "count": 2, "value": 300}],
            "vacations": [
                {
                    "employee": "Иван Иванов",
                    "department": "Отдел разработки",
                    "type": "ANNUAL",
                    "startDate": "2026-07-01",
                    "endDate": "2026-07-14",
                }
            ],
        },
        "activeProjects": [
            {
                "code": "PR-7",
                "name": "Проект Север",
                "endDate": "2026-12-01",
                "progress": 70,
                "tasksTotal": 10,
                "tasksCompleted": 7,
                "plannedBudget": 5000,
                "actualBudget": 4500,
            }
        ],
        "overdueTasks": [
            {
                "name": "Подготовить релиз",
                "project": {"code": "PR-7"},
                "responsible": "Иван Иванов",
                "overdueDays": 4,
            }
        ],
        "contracts": [
            {
                "employee": "Иван Иванов",
                "department": "Отдел разработки",
                "contractEndDate": "2026-11-01",
            }
        ],
        "assetAttention": [
            {"name": "Ноутбук", "inventoryNumber": "INV-7", "status": "UNDER_REPAIR"}
        ],
    }
    columns, rows = _dashboard_report_rows(data)
    docx = Document(BytesIO(render_docx("Отчёт", columns, rows)))
    docx_text = "\n".join(cell.text for row in docx.tables[0].rows for cell in row.cells)
    assert all(
        value in docx_text
        for value in (
            "Отдел разработки",
            "Проект Север",
            "Ноутбук",
            "Иван Иванов",
            "Подготовить релиз",
            "1000",
        )
    )

    pdf = render_pdf("Отчёт", columns, rows)
    assert pdf.startswith(b"%PDF-")
    pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(pdf)).pages)
    assert all(
        value in pdf_text
        for value in ("Проект Север", "INV-7", "Отдел разработки", "Подготовить релиз", "1000")
    )


def test_all_dashboard_export_formats_hide_finance_without_finance_permission() -> None:
    private_salary = "987654321.25"
    data = {
        "query": {"dateFrom": "2026-01-01", "dateTo": "2026-12-31"},
        "visibility": {"employees": True, "finance": False, "projects": False, "assets": False},
        "summary": {
            "headcount": 1,
            "active": 1,
            "occupiedRate": 1,
            "staffRate": 1,
            "plannedFot": private_salary,
            "actualFot": private_salary,
        },
        "analytics": {
            "workforceByDepartment": [
                {
                    "department": "Закрытый отдел",
                    "headcount": 1,
                    "active": 1,
                    "occupiedRate": 1,
                    "staffRate": 1,
                    "occupancy": 100,
                    "plannedFot": private_salary,
                    "actualFot": private_salary,
                }
            ],
            "assetsByStatus": [],
            "vacations": [],
        },
        "activeProjects": [],
        "overdueTasks": [],
        "contracts": [],
        "assetAttention": [],
    }

    workbook = load_workbook(BytesIO(render_dashboard_xlsx(data)), data_only=True)
    workbook_values: list[Any] = []
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        assert sheet is not None
        workbook_values.extend(value for row in sheet.values for value in row)
    workforce = workbook["Сотрудники"]
    assert "Плановый ФОТ" not in [cell.value for cell in workforce[1]]
    assert "Фактический ФОТ" not in [cell.value for cell in workforce[1]]
    assert private_salary not in [str(value) for value in workbook_values]

    columns, rows = _dashboard_report_rows(data)
    docx = Document(BytesIO(render_docx("Отчёт", columns, rows)))
    docx_text = "\n".join(cell.text for row in docx.tables[0].rows for cell in row.cells)
    pdf = render_pdf("Отчёт", columns, rows)
    pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(pdf)).pages)
    for text in (docx_text, pdf_text):
        assert private_salary not in text
        assert "Плановый ФОТ" not in text
        assert "Фактический ФОТ" not in text
