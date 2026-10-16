from __future__ import annotations

from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.domain.assets.importing import (
    ASSET_IMPORT_COLUMNS,
    MAX_ASSET_IMPORT_FILE_BYTES,
    MAX_ASSET_IMPORT_ROWS,
    AssetWorkbookFormatError,
)
from app.integrations.asset_import import (
    OpenpyxlAssetWorkbook,
    build_asset_import_template,
    build_export_workbook,
    parse_asset_import_workbook,
)
from app.services.asset_import import AssetImportError, AssetImportService


def _xlsx_bytes(workbook: Workbook) -> bytes:
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _import_workbook(headers: list[str] | None = None) -> tuple[Workbook, object]:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Имущество"
    sheet.append(headers or [column.label for column in ASSET_IMPORT_COLUMNS])
    return workbook, sheet


def _valid_values() -> dict[str, object]:
    return {
        "inventoryNumber": " INV-1 ",
        "name": "Ноутбук",
        "unitPrice": 1250.5,
        "unitOfMeasure": "шт",
        "quantity": 2,
        "molCode": "M-1",
        "groupCode": "G-1",
        "recordingDate": datetime(2026, 7, 16),
        "documentType": "Накладная",
        "documentDetails": "№ 7",
        "contractCode": "",
        "internalFundingCode": "",
        "isExistingAsset": "Нет",
        "accountingForm": "145",
        "notes": "",
    }


def test_parser_maps_reordered_headers_and_skips_blank_rows() -> None:
    workbook, sheet = _import_workbook()
    columns = [column.label for column in reversed(ASSET_IMPORT_COLUMNS)]
    sheet.delete_rows(1)
    sheet.append(columns)
    values = _valid_values()
    sheet.append([values[column.key] for column in reversed(ASSET_IMPORT_COLUMNS)])
    sheet.append([None] * len(ASSET_IMPORT_COLUMNS))
    values["inventoryNumber"] = "INV-2"
    sheet.append([values[column.key] for column in reversed(ASSET_IMPORT_COLUMNS)])

    rows = parse_asset_import_workbook(_xlsx_bytes(workbook))

    assert [row.rowNumber for row in rows] == [2, 4]
    assert rows[0].values["inventoryNumber"] == " INV-1 "
    assert rows[0].values["unitPrice"] == 1250.5
    assert rows[0].values["recordingDate"] == datetime(2026, 7, 16)


@pytest.mark.parametrize(
    ("headers", "message"),
    [
        (["Наименование"], "Не найдены обязательные колонки"),
        (
            [column.label for column in ASSET_IMPORT_COLUMNS] + ["Наименование"],
            "указана дважды",
        ),
        (
            [column.label for column in ASSET_IMPORT_COLUMNS] + ["Неизвестная колонка"],
            "Неизвестные колонки",
        ),
    ],
)
def test_parser_rejects_missing_duplicate_and_unknown_headers(
    headers: list[str], message: str
) -> None:
    workbook, _ = _import_workbook(headers)

    with pytest.raises(AssetWorkbookFormatError, match=message):
        parse_asset_import_workbook(_xlsx_bytes(workbook))


def test_parser_rejects_missing_sheet_empty_data_and_corrupt_zip() -> None:
    workbook = Workbook()
    workbook.active.title = "Другая вкладка"
    with pytest.raises(AssetWorkbookFormatError, match="нет листа"):
        parse_asset_import_workbook(_xlsx_bytes(workbook))

    empty_workbook, _ = _import_workbook()
    with pytest.raises(AssetWorkbookFormatError, match="нет строк с имуществом"):
        parse_asset_import_workbook(_xlsx_bytes(empty_workbook))

    with pytest.raises(AssetWorkbookFormatError, match="корректный шаблон"):
        parse_asset_import_workbook(b"not a zip workbook")


def test_parser_rejects_more_than_one_hundred_data_rows() -> None:
    workbook, sheet = _import_workbook()
    for index in range(MAX_ASSET_IMPORT_ROWS + 1):
        row = [None] * len(ASSET_IMPORT_COLUMNS)
        row[0] = f"INV-{index}"
        sheet.append(row)

    with pytest.raises(AssetWorkbookFormatError, match="больше 100 строк"):
        parse_asset_import_workbook(_xlsx_bytes(workbook))


def test_import_service_rejects_file_over_five_megabytes() -> None:
    with pytest.raises(AssetImportError, match="Максимальный размер — 5 МБ"):
        AssetImportService(lambda: None, OpenpyxlAssetWorkbook()).parse_workbook(  # type: ignore[arg-type]
            b"x" * (MAX_ASSET_IMPORT_FILE_BYTES + 1)
        )


def test_import_template_contains_three_sheets_scoped_lookups_and_dropdowns() -> None:
    payload = build_asset_import_template(
        [{"code": "M-1", "fullName": "Иванов И.И.", "department": "ИТ"}],
        [{"code": "G-1", "name": "Оргтехника"}],
    )
    workbook = load_workbook(BytesIO(payload), data_only=False)
    assert workbook.sheetnames == ["Как заполнить", "Имущество", "Справочники"]
    assert workbook["Имущество"]["A1"].value == "Инвентарный номер"
    assert workbook["Имущество"].freeze_panes == "A2"
    assert workbook["Имущество"].auto_filter.ref == "A1:O101"
    assert workbook["Имущество"]["C2"].number_format == "#,##0.00"
    assert workbook["Имущество"]["H2"].number_format == "dd.mm.yyyy"
    assert workbook["Справочники"]["A2"].value == "M-1"
    assert workbook["Справочники"]["E2"].value == "Оргтехника"
    assert workbook.defined_names["MolCodes"].attr_text == "'Справочники'!$A$2:$A$2"
    assert workbook.defined_names["AssetGroupCodes"].attr_text == "'Справочники'!$D$2:$D$2"
    formulas = {rule.formula1 for rule in workbook["Имущество"].data_validations.dataValidation}
    assert "=MolCodes" in formulas
    assert "=AssetGroupCodes" in formulas
    assert '"Да,Нет"' in formulas
    assert '"145,367"' in formulas


def test_export_workbook_has_data_sheet_russian_headers_filters_and_numeric_cells() -> None:
    payload = build_export_workbook(
        ["Инвентарный номер", "Количество", "Статус", "Дата", "Примечание"],
        [
            {
                "Инвентарный номер": "INV-1",
                "Количество": 2.5,
                "Статус": "В наличии",
                "Дата": "16.07.2026",
                "Примечание": '=HYPERLINK("https://example.invalid")',
            }
        ],
    )
    workbook = load_workbook(BytesIO(payload), data_only=True)
    assert workbook.sheetnames == ["Data"]
    sheet = workbook["Data"]
    assert list(sheet.values) == [
        ("Инвентарный номер", "Количество", "Статус", "Дата", "Примечание"),
        ("INV-1", 2.5, "В наличии", "16.07.2026", '=HYPERLINK("https://example.invalid")'),
    ]
    assert sheet["B2"].data_type == "n"
    assert sheet["E2"].data_type == "s"
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:E2"
