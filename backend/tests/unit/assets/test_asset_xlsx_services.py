from __future__ import annotations

from contextlib import AbstractAsyncContextManager
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.core.context import Actor, RequestContext
from app.domain.assets.importing import ParsedAssetImportRow
from app.domain.auth.rules import AccessContext, ScopeGrant
from app.integrations.asset_import import OpenpyxlAssetWorkbook
from app.services.asset_export import AssetExportService
from app.services.asset_import import AssetImportService


class _Assets:
    def __init__(self) -> None:
        self.mols = [
            {
                "id": "mol-1",
                "code": "M-1",
                "departmentId": "department-1",
                "fullName": "Иванов И.И.",
                "department": "ИТ",
                "storageLocation": "Кабинет 1",
            }
        ]
        self.existing: set[str] = set()
        self.assets: tuple[object, ...] = ()
        self.operations: list[dict[str, object]] = []

    async def mol_options(self, allowed: tuple[str, ...] | None) -> list[dict[str, object]]:
        return self.mols

    async def existing_numbers(self, numbers: list[str]) -> set[str]:
        return self.existing.intersection(numbers)

    async def export_assets(
        self, archived: bool | None, access: AccessContext
    ) -> tuple[object, ...]:
        return self.assets

    async def export_operations(self, access: AccessContext) -> list[dict[str, object]]:
        return self.operations


class _Groups:
    def __init__(self) -> None:
        self.groups = [SimpleNamespace(id="group-1", code="G-1", name="Оргтехника")]

    async def list(self) -> list[SimpleNamespace]:
        return self.groups


class _Uow(AbstractAsyncContextManager["_Uow"]):
    def __init__(self) -> None:
        self.assets = _Assets()
        self.groups = _Groups()

    async def __aenter__(self) -> _Uow:
        return self

    async def __aexit__(self, *args: object) -> None:
        return None


def _factory(uow: _Uow):
    return lambda: uow


def _context() -> RequestContext:
    access = AccessContext(
        "user-1",
        None,
        None,
        (),
        (ScopeGrant("assignment-1", "ADMIN", "ALL", "ALL"),),
    )
    actor = Actor("user-1", is_authenticated=True, access=access)
    return RequestContext("request-1", actor=actor)


def _row(row_number: int, **overrides: object) -> ParsedAssetImportRow:
    values: dict[str, object] = {
        "inventoryNumber": "INV-1",
        "name": " Ноутбук ",
        "unitPrice": "1 250,50",
        "unitOfMeasure": "шт",
        "quantity": 2,
        "molCode": "M-1",
        "groupCode": "G-1",
        "recordingDate": "16.07.2026",
        "documentType": "Накладная",
        "documentDetails": "№ 7",
        "isExistingAsset": "Нет",
        "accountingForm": "145",
    }
    values.update(overrides)
    return ParsedAssetImportRow(row_number, values)


@pytest.mark.anyio
async def test_preview_normalizes_values_uses_exact_codes_and_marks_duplicates() -> None:
    uow = _Uow()
    uow.assets.existing = {"INV-DB"}
    service = AssetImportService(_factory(uow), OpenpyxlAssetWorkbook())  # type: ignore[arg-type]

    preview = await service.preview(
        (
            _row(2, inventoryNumber="INV-VALID"),
            _row(3),
            _row(4),
            _row(5, inventoryNumber="INV-DB"),
            _row(6, molCode="m-1"),
            _row(7, unitPrice="12,345", recordingDate="31.02.2026"),
        ),
        _context(),
    )

    assert preview[0]["input"] == {
        "inventoryNumber": "INV-VALID",
        "name": "Ноутбук",
        "unitPrice": 1250.5,
        "unitOfMeasure": "шт",
        "quantity": 2.0,
        "molId": "mol-1",
        "groupId": "group-1",
        "recordingDate": "2026-07-16",
        "documentType": "Накладная",
        "documentDetails": "№ 7",
        "isExistingAsset": False,
        "accountingForm": "145",
        "projectId": None,
        "status": "IN_STOCK",
        "documentFiles": [],
        "photos": [],
    }
    assert preview[1]["errors"] == ["Инвентарный номер повторяется в этом файле"]
    assert preview[2]["errors"] == ["Инвентарный номер повторяется в этом файле"]
    assert preview[3]["errors"] == ["Инвентарный номер уже используется"]
    assert "МОЛ с таким кодом не найден или недоступен" in preview[4]["errors"]
    assert "Цена за единицу: допустимо не более двух знаков после запятой" in preview[5]["errors"]
    assert "Дата поступления: Дата поступления обязательна" in preview[5]["errors"]


@pytest.mark.anyio
async def test_export_service_keeps_legacy_headers_labels_and_numeric_types() -> None:
    uow = _Uow()
    uow.assets.assets = (
        SimpleNamespace(
            id="asset-1",
            orderNumber=7,
            name="Ноутбук",
            inventoryNumber="INV-1",
            group={"name": "Оргтехника"},
            holdings=(
                SimpleNamespace(
                    mol={
                        "fullName": "Иванов И.И.",
                        "department": "ИТ",
                        "storageLocation": "Кабинет 1",
                    },
                    quantity=Decimal("2.00"),
                ),
            ),
            unitPrice=Decimal("1250.50"),
            unitOfMeasure="шт",
            quantity=Decimal("2.00"),
            totalCost=Decimal("2501.00"),
            status="IN_STOCK",
            recordingDate=datetime(2026, 7, 16),
            documentType="Накладная",
            documentDetails="№ 7",
            contractCode=None,
            internalFundingCode=None,
            isExistingAsset=False,
            isArchived=False,
            plannedDisposalDate=None,
            plannedDisposalReason=None,
            notes=None,
            documentFiles=[],
            photos=[],
            createdAt=datetime(2026, 7, 16),
            updatedAt=datetime(2026, 7, 16),
        ),
    )
    service = AssetExportService(_factory(uow), OpenpyxlAssetWorkbook())  # type: ignore[arg-type]

    payload, filename = await service.export("assets", _context().actor, {})  # type: ignore[arg-type]

    workbook = load_workbook(BytesIO(payload), data_only=True)
    sheet = workbook["Data"]
    assert filename.startswith("assets_export_") and filename.endswith(".xlsx")
    assert sheet["A1"].value == "ID"
    assert sheet["D2"].value == "INV-1"
    assert sheet["F2"].value == "Иванов И.И. (2.00)"
    assert sheet["G2"].value == "ИТ"
    assert sheet["I2"].value == 1250.5
    assert sheet["I2"].data_type == "n"
    assert sheet["M2"].value == "В наличии"
    assert sheet["N2"].value == "16.07.2026"
    assert sheet.freeze_panes == "A2"


@pytest.mark.anyio
async def test_export_service_keeps_operations_headers_and_labels() -> None:
    uow = _Uow()
    uow.assets.operations = [
        {
            "id": "operation-1",
            "type": "TRANSFER",
            "asset": {
                "name": "Ноутбук",
                "inventoryNumber": "INV-1",
                "group": {"name": "Оргтехника"},
            },
            "fromMol": {"fullName": "Иванов И.И."},
            "toMol": {"fullName": "Петров П.П."},
            "quantity": Decimal("1.00"),
            "unitPrice": Decimal("1250.50"),
            "totalCost": Decimal("1250.50"),
            "date": datetime(2026, 7, 17),
            "reason": None,
            "documentType": "Акт",
            "documentDetails": "№ 8",
            "oldStatus": "IN_STOCK",
            "newStatus": "IN_USE",
            "createdAt": datetime(2026, 7, 17),
        }
    ]
    service = AssetExportService(_factory(uow), OpenpyxlAssetWorkbook())  # type: ignore[arg-type]

    payload, filename = await service.export("operations", _context().actor, {})  # type: ignore[arg-type]

    sheet = load_workbook(BytesIO(payload), data_only=True)["Data"]
    assert filename.startswith("operations_export_") and filename.endswith(".xlsx")
    assert sheet["A1"].value == "ID операции"
    assert sheet["B2"].value == "Передача"
    assert sheet["D2"].value == "INV-1"
    assert sheet["F2"].value == "Иванов И.И."
    assert sheet["G2"].value == "Петров П.П."
    assert sheet["O2"].value == "В наличии"
    assert sheet["P2"].value == "В эксплуатации"
    assert sheet["K2"].value == "17.07.2026"


@pytest.mark.anyio
async def test_template_service_queries_scoped_mols_and_returns_the_binary_template() -> None:
    uow = _Uow()
    service = AssetImportService(_factory(uow), OpenpyxlAssetWorkbook())  # type: ignore[arg-type]

    payload = await service.template(_context())

    workbook = load_workbook(BytesIO(payload), data_only=False)
    assert workbook["Справочники"]["A2"].value == "M-1"
    assert workbook["Справочники"]["B2"].value == "Иванов И.И."
    assert workbook["Справочники"]["C2"].value == "ИТ"
    assert workbook["Справочники"]["D2"].value == "G-1"
